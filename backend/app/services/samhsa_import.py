"""CSV / Excel template + bulk import for SAMHSA-seeded rehab listings."""
from __future__ import annotations

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Iterable

ProgressFn = Callable[..., None]

from sqlalchemy.orm import Session

from app.models.rehab import CenterSource, ListingStatus, RehabCenter
from app.models.service_code import ServiceCodeCatalog
from app.services.listing_media import LISTING_PLACEHOLDER_IMAGE, is_custom_listing_image
from app.services.service_codes import (
    catalog_lookup,
    derive_listing_fields,
    resolve_imported_service_codes,
    seed_service_code_catalog,
)

CSV_HEADERS = [
    "name1",
    "name2",
    "street1",
    "street2",
    "city",
    "state",
    "zip",
    "phone",
    "intake1",
    "intake2",
    "intake1a",
    "intake2a",
    "service_code_info",
]

# Common SAMHSA locator / spreadsheet aliases → canonical template keys.
HEADER_ALIASES = {
    "samhsa_id": ("samhsa_id", "samhsaid", "npi", "facility_id", "locator_id"),
    "name": ("name", "facility_name", "facility", "center_name", "listing_name", "center"),
    "name1": ("name1", "business_name"),
    "name2": ("name2", "title", "listing_title", "program_name"),
    "address_line": ("address_line", "address", "street", "street_address"),
    "street1": ("street1", "address1"),
    "street2": ("street2", "address2", "address_line2", "address_line_2", "street_address_2", "suite"),
    "city": ("city",),
    "state": ("state", "state_province", "st"),
    "zip": ("zip", "zipcode", "zip_code", "postal_code"),
    "phone": ("phone", "telephone", "phone_number", "phone1"),
    "intake1": ("intake1", "intake_1"),
    "intake2": ("intake2", "intake_2"),
    "intake1a": ("intake1a", "intake_1a", "intake1_a"),
    "intake2a": ("intake2a", "intake_2a", "intake2_a"),
    "website": ("website", "url", "website_url", "web"),
    "outreach_email": ("outreach_email", "email", "e_mail", "facility_email"),
    "contact_email": ("contact_email", "admissions_email", "info_email"),
    "description": ("description", "about", "summary"),
    "specialties": ("specialties", "services", "service_types"),
    "levels_of_care": ("levels_of_care", "level_of_care", "care_levels"),
    "service_codes": ("service_codes", "service_code", "samhsa_codes", "svc_codes"),
    "service_code_info": (
        "service_code_info",
        "service_codes_info",
        "service_code_information",
        "samhsa_service_code_info",
        "locator_service_codes",
        "svc_info",
        "serviceinfo",
        "service_info",
    ),
    "insurances": ("insurances", "insurance", "accepted_insurance"),
    "amenities": ("amenities",),
    "accreditations": ("accreditations", "accreditation"),
    "google_maps_url": ("google_maps_url", "maps_url", "map_url"),
    "google_reviews_url": ("google_reviews_url", "reviews_url"),
    "rating": ("rating", "stars"),
}

IMAGE_COLUMNS = {
    "image", "image_url", "image_key", "photo", "photos", "picture", "thumbnail",
    "gallery", "gallery_urls", "gallery_keys", "logo", "logo_url", "logo_key",
}

ALIAS_TO_CANONICAL = {
    alias: canonical
    for canonical, aliases in HEADER_ALIASES.items()
    for alias in aliases
}

TEMPLATE_EXAMPLE_ROW = {
    "name1": "Example Recovery LLC",
    "name2": "Example Recovery Center",
    "street1": "123 Main St",
    "street2": "Suite 100",
    "city": "Phoenix",
    "state": "Arizona",
    "zip": "85001",
    "phone": "602-555-0100",
    "intake1": "602-555-0100",
    "intake2": "800-555-0199",
    "intake1a": "",
    "intake2a": "",
    "service_code_info": "SA MH SUMH * OP * CMHC * CHLOR FLUPH HALOP ARIPI ASENA BREXP CARIP CLOZA ILOPE LURAS OLANZ PALIP QUETI RISPE ZIPRA ANTPYCH * CBT CFT GT IDD IPT TELE * CIT WI * PVTN * CLF CMHG MC MD TRICARE OSF PI SCJJ CASH SMHA SWFS VAF * SS * YA SEN VET ADMIL MILF CJ MHSU HIV TRMA PTSD SED SMI * STU * SMPD * CH/AD YAD ADLT SNR * VPPD * ACT COOT FPSY PRS CM SPS",
}

MAX_IMPORT_ROWS = 25_000
MAX_IMPORT_BYTES = 40 * 1024 * 1024
INSERT_BATCH = 500
ERROR_CAP = 80
UPDATE_SKIP = {"source"}
_FLAG_OFF = {"", "0", "n", "no", "false", "f", "off", "none", "null", "-", "na"}


def _slugify(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", (name or "").lower()).strip("-")
    return s[:200] or "center"


def _listing_name(row: dict[str, str]) -> str:
    """name2 is the public title; fall back to name1 (business) or a generic name column."""
    return _clean(row.get("name2")) or _clean(row.get("name1")) or _clean(row.get("name")) or ""


def _listing_address(row: dict[str, str]) -> str | None:
    """street1 is address1; street2 is address2 (suite / unit). Combine when both are present."""
    line1 = _clean(row.get("street1")) or _clean(row.get("address_line"))
    line2 = _clean(row.get("street2"))
    if line1 and line2:
        return f"{line1}, {line2}"[:255]
    return line1 or line2


def _listing_phone(row: dict[str, str]) -> str | None:
    """Use the phone column, or the first available SAMHSA intake number."""
    for key in ("phone", "intake1", "intake2", "intake1a", "intake2a"):
        value = _clean(row.get(key))
        if value:
            return value[:50]
    return None


def _split_list(value: str | None) -> list[str]:
    if not value or not str(value).strip():
        return []
    parts = re.split(r"[|;,]", str(value))
    return [p.strip() for p in parts if p.strip()]


def _code_header_keys(code: str) -> set[str]:
    raw = (code or "").strip()
    variants = (
        raw,
        raw.replace("/", "_"),
        raw.replace("/", ""),
        raw.replace("-", "_"),
        raw.replace(" ", "_"),
    )
    return {_normalize_header(item).upper() for item in variants if item}


def _flag_header_map(headers: list[str], service_lookup: dict) -> dict[str, str]:
    """Map spreadsheet headers that are catalog codes (SA, OP, CH/AD, …) to the stored code."""
    by_header: dict[str, str] = {}
    for code, row in service_lookup.items():
        for key in _code_header_keys(code):
            by_header[key] = row.service_code
    mapped: dict[str, str] = {}
    for header in headers:
        normalized = _normalize_header(header)
        if not normalized or ALIAS_TO_CANONICAL.get(normalized):
            continue
        code = by_header.get(normalized.upper())
        if code:
            mapped[header] = code
    return mapped


def _is_flag_on(value) -> bool:
    return _cell_str(value).strip().lower() not in _FLAG_OFF


def _codes_from_flag_columns(raw: dict, header_map: dict[str, str]) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()
    for header, code in header_map.items():
        if not _is_flag_on(raw.get(header)):
            continue
        if code in seen:
            continue
        seen.add(code)
        found.append(code)
    return found


def _import_service_codes(
    row: dict[str, str],
    service_lookup: dict | None,
    raw: dict | None = None,
    flag_map: dict[str, str] | None = None,
) -> list[str]:
    chunks = [row.get("service_code_info") or "", row.get("service_codes") or ""]
    combined = " * ".join(chunk for chunk in chunks if chunk.strip())
    codes: list[str] = []
    if service_lookup:
        codes = [item.service_code for item in resolve_imported_service_codes(combined, service_lookup)]
    elif combined:
        codes = _split_list(combined) or [token for token in re.split(r"[\s*|;,]+", combined) if token.strip()]
    if raw and flag_map:
        for code in _codes_from_flag_columns(raw, flag_map):
            if code not in codes:
                codes.append(code)
    return codes


def _fill_from_service_codes(fields: dict, service_lookup: dict | None) -> None:
    if not service_lookup or not fields.get("service_codes"):
        return
    rows = [service_lookup[code.upper()] for code in fields["service_codes"] if code.upper() in service_lookup]
    derived = derive_listing_fields(rows)
    if not fields.get("specialties") and derived["specialties"]:
        fields["specialties"] = derived["specialties"]
    if not fields.get("levels_of_care") and derived["levels_of_care"]:
        fields["levels_of_care"] = derived["levels_of_care"]
    if not fields.get("insurances") and derived["insurances"]:
        fields["insurances"] = derived["insurances"]


def _clean(value: str | None) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def _location_display(city: str | None, state: str | None) -> str:
    parts = [p for p in (city, state) if p]
    return ", ".join(parts)


def _normalize_header(raw: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", "_", (raw or "").strip().lower()).strip("_")


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).rstrip("0").rstrip(".")
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def _canonical_row(raw: dict) -> dict[str, str]:
    row: dict[str, str] = {}
    for key, value in raw.items():
        normalized = _normalize_header(key)
        if normalized in IMAGE_COLUMNS:
            continue
        canonical = ALIAS_TO_CANONICAL.get(normalized)
        if not canonical or canonical in row:
            continue
        text = _cell_str(value)
        if canonical == "zip" and text.isdigit() and 3 <= len(text) <= 4:
            text = text.zfill(5)
        row[canonical] = text
    return row


def _detect_dialect(text: str) -> type[csv.Dialect] | csv.Dialect:
    sample = text[:8192]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        return csv.excel


def _file_kind(content: bytes, filename: str = "") -> str:
    name = (filename or "").lower()
    if content.startswith(b"PK"):
        return "xlsx"
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith(".xls"):
        return "xls"
    return "csv"


def _emit(on_progress: ProgressFn | None, **payload) -> None:
    if on_progress:
        on_progress(**payload)


def _rows_from_csv(content: bytes) -> tuple[list[str], list[dict]]:
    text = content.decode("utf-8-sig")
    if text.startswith("\ufeff"):
        text = text.lstrip("\ufeff")
    dialect = _detect_dialect(text)
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames:
        return [], []
    headers = [h for h in reader.fieldnames if h]
    rows = [{k: v for k, v in raw.items() if k} for raw in reader]
    return headers, rows


def _rows_from_xlsx(content: bytes, on_progress: ProgressFn | None = None) -> tuple[list[str], list[dict]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel .xlsx support is missing. Install openpyxl.") from exc
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration:
            return [], []
        headers = [_cell_str(v) for v in header_row]
        rows: list[dict] = []
        for n, values in enumerate(iterator, start=1):
            raw = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers)) if headers[i]}
            rows.append(raw)
            if n % 500 == 0:
                _emit(
                    on_progress,
                    phase="reading",
                    message=f"Reading spreadsheet… {n:,} rows",
                    processed=n,
                    percent=min(10, 3 + n // 2000),
                )
        return [h for h in headers if h], rows
    finally:
        wb.close()


def _rows_from_xls(content: bytes, on_progress: ProgressFn | None = None) -> tuple[list[str], list[dict]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("Excel .xls support is missing. Install xlrd.") from exc
    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    if sheet.nrows < 1:
        return [], []
    headers = [_cell_str(sheet.cell_value(0, c)) for c in range(sheet.ncols)]
    rows: list[dict] = []
    total = max(sheet.nrows - 1, 0)
    for r in range(1, sheet.nrows):
        raw = {headers[c]: sheet.cell_value(r, c) for c in range(sheet.ncols) if headers[c]}
        rows.append(raw)
        if r % 500 == 0:
            _emit(
                on_progress,
                phase="reading",
                message=f"Reading spreadsheet… {r:,} of {total:,} rows",
                processed=r,
                total=total,
                percent=min(10, 3 + int(r / max(total, 1) * 7)),
            )
    return [h for h in headers if h], rows


def parse_import_table(
    content: str | bytes,
    filename: str = "",
    on_progress: ProgressFn | None = None,
) -> tuple[list[str], list[dict]]:
    if isinstance(content, str):
        payload = content.encode("utf-8")
    else:
        payload = content
    kind = _file_kind(payload, filename)
    if kind == "xlsx":
        return _rows_from_xlsx(payload, on_progress)
    if kind == "xls":
        return _rows_from_xls(payload, on_progress)
    return _rows_from_csv(payload)


def _unique_slug(db: Session, base: str, exclude_id: int | None = None) -> str:
    slug = base
    n = 1
    while True:
        q = db.query(RehabCenter).filter(RehabCenter.slug == slug)
        if exclude_id is not None:
            q = q.filter(RehabCenter.id != exclude_id)
        if not q.first():
            return slug
        slug = f"{base}-{n}"
        n += 1


def _next_slug(base: str, taken: set[str]) -> str:
    slug = base
    n = 1
    while slug in taken:
        slug = f"{base}-{n}"
        n += 1
    taken.add(slug)
    return slug


def _match_key(name: str | None, city: str | None, state: str | None) -> tuple[str, str, str]:
    return (
        (name or "").strip().lower(),
        (city or "").strip().lower(),
        (state or "").strip().lower(),
    )


def build_template_csv() -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=CSV_HEADERS, lineterminator="\n")
    writer.writeheader()
    writer.writerow(TEMPLATE_EXAMPLE_ROW)
    return buf.getvalue()


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)
    total_rows: int = 0


def _find_existing(db: Session, samhsa_id: str | None, name: str, city: str | None, state: str | None) -> RehabCenter | None:
    if samhsa_id:
        found = db.query(RehabCenter).filter(RehabCenter.samhsa_id == samhsa_id).first()
        if found:
            return found
    q = db.query(RehabCenter).filter(RehabCenter.name.ilike(name.strip()))
    if city:
        q = q.filter(RehabCenter.city.ilike(city.strip()))
    if state:
        q = q.filter(RehabCenter.state.ilike(state.strip()))
    return q.first()


def _row_fields(
    row: dict[str, str],
    name: str,
    *,
    rating: float,
    service_lookup: dict | None = None,
    raw: dict | None = None,
    flag_map: dict[str, str] | None = None,
) -> dict:
    city = _clean(row.get("city"))
    state = _clean(row.get("state"))
    payload = {
        "name": name,
        "address_line": _listing_address(row),
        "city": city,
        "state": state,
        "zip": _clean(row.get("zip")),
        "phone": _listing_phone(row),
        "intake1": _clean(row.get("intake1")),
        "intake2": _clean(row.get("intake2")),
        "intake1a": _clean(row.get("intake1a")),
        "intake2a": _clean(row.get("intake2a")),
        "website": _clean(row.get("website")),
        "outreach_email": _clean(row.get("outreach_email")),
        "contact_email": _clean(row.get("contact_email")),
        "description": row.get("description") or "",
        "specialties": _split_list(row.get("specialties")),
        "levels_of_care": _split_list(row.get("levels_of_care")),
        "service_codes": _import_service_codes(row, service_lookup, raw=raw, flag_map=flag_map),
        "insurances": _split_list(row.get("insurances")),
        "amenities": _split_list(row.get("amenities")),
        "accreditations": _split_list(row.get("accreditations")),
        "google_maps_url": _clean(row.get("google_maps_url")),
        "google_reviews_url": _clean(row.get("google_reviews_url")),
        "rating": rating,
        "location_display": _location_display(city, state),
        "samhsa_id": _clean(row.get("samhsa_id")),
    }
    _fill_from_service_codes(payload, service_lookup)
    return payload


def _note(result: ImportResult, message: str) -> None:
    if len(result.errors) < ERROR_CAP:
        result.errors.append(message)


def _parse_rating(raw: str | None, result: ImportResult, row_no: int) -> float:
    rating_raw = _clean(raw)
    if not rating_raw:
        return 5.0
    try:
        rating = float(rating_raw)
    except ValueError:
        _note(result, f"Row {row_no}: invalid rating — defaulted to 5.0")
        return 5.0
    if rating < 0 or rating > 5:
        return 5.0
    return rating


def _apply_update(existing: RehabCenter, fields: dict, *, publish: bool) -> None:
    for key, value in fields.items():
        if key in UPDATE_SKIP:
            continue
        if key == "samhsa_id" and existing.samhsa_id and fields.get("samhsa_id") and existing.samhsa_id != fields["samhsa_id"]:
            continue
        setattr(existing, key, value)
    if not is_custom_listing_image(existing.image_key):
        existing.image_key = LISTING_PLACEHOLDER_IMAGE
    if not existing.gallery_keys:
        existing.gallery_keys = []
    if publish and existing.listing_status == ListingStatus.draft and not existing.claimed:
        existing.listing_status = ListingStatus.published
        if not existing.published_at:
            existing.published_at = datetime.now(timezone.utc)


def _new_mapping(fields: dict, slug: str, *, publish: bool, now: datetime) -> dict:
    return {
        **fields,
        "slug": slug,
        "listing_status": ListingStatus.published if publish else ListingStatus.draft,
        "published_at": now if publish else None,
        "claimed": False,
        "contact_visible": False,
        "image_key": LISTING_PLACEHOLDER_IMAGE,
        "gallery_keys": [],
        "testimonials": [],
        "source": CenterSource.imported,
        "inquiry_form_enabled": True,
    }


def _chunked(items: list, size: int) -> Iterable[list]:
    for i in range(0, len(items), size):
        yield items[i:i + size]


def _import_percent(phase: str, processed: int, total: int) -> int:
    if phase == "done":
        return 100
    if phase == "error":
        return 100
    if phase == "reading":
        return 8
    if phase == "importing" and total > 0:
        return 12 + int(processed / total * 73)
    if phase == "saving" and total > 0:
        return 85 + int(processed / total * 13)
    if phase == "saving":
        return 90
    return 4


def import_centers_file(
    db: Session,
    content: str | bytes,
    *,
    filename: str = "import.csv",
    publish: bool = True,
    on_progress: ProgressFn | None = None,
) -> ImportResult:
    if isinstance(content, bytes) and len(content) > MAX_IMPORT_BYTES:
        return ImportResult(errors=[f"File is too large. Upload a CSV or Excel file under {MAX_IMPORT_BYTES // (1024 * 1024)} MB."])

    _emit(on_progress, phase="reading", message="Reading file…", percent=4)
    try:
        headers, raw_rows = parse_import_table(content, filename, on_progress)
    except RuntimeError as exc:
        return ImportResult(errors=[str(exc)])
    except Exception as exc:  # noqa: BLE001
        return ImportResult(errors=[f"Could not read this file: {exc}"])

    if not headers:
        return ImportResult(errors=["File is empty or missing a header row."])

    mapped = {ALIAS_TO_CANONICAL.get(_normalize_header(h)) for h in headers if _normalize_header(h) in ALIAS_TO_CANONICAL}
    if not mapped.intersection({"name", "name1", "name2"}):
        return ImportResult(errors=["Missing required column: name2 (title), name1 (business name), or name."])

    if len(raw_rows) > MAX_IMPORT_ROWS:
        return ImportResult(errors=[f"File has {len(raw_rows):,} rows. Import up to {MAX_IMPORT_ROWS:,} listings at a time."])

    _emit(
        on_progress,
        phase="reading",
        message=f"Preparing {len(raw_rows):,} rows…",
        processed=0,
        total=len(raw_rows),
        percent=10,
    )
    existing_by_samhsa: dict[str, RehabCenter] = {}
    existing_by_key: dict[tuple[str, str, str], RehabCenter] = {}
    taken_slugs: set[str] = set()
    for center in db.query(RehabCenter).all():
        if center.slug:
            taken_slugs.add(center.slug)
        if center.deleted_at:
            continue
        if center.samhsa_id:
            existing_by_samhsa[center.samhsa_id] = center
        existing_by_key[_match_key(center.name, center.city, center.state)] = center

    if db.query(ServiceCodeCatalog).count() == 0:
        seed_service_code_catalog(db)
    service_lookup = catalog_lookup(db)
    flag_map = _flag_header_map(headers, service_lookup)

    result = ImportResult()
    inserts: list[dict] = []
    pending_by_samhsa: dict[str, dict] = {}
    pending_by_key: dict[tuple[str, str, str], dict] = {}
    now = datetime.now(timezone.utc)
    planned = len(raw_rows)

    for i, raw in enumerate(raw_rows, start=2):
        row = _canonical_row(raw)
        if not any(row.values()):
            continue
        result.total_rows += 1
        name = _listing_name(row)
        if not name:
            result.skipped += 1
            _note(result, f"Row {i}: name2 (title) or name1 (business name) is required — skipped")
            continue

        fields = _row_fields(
            row,
            name,
            rating=_parse_rating(row.get("rating"), result, i),
            service_lookup=service_lookup,
            raw=raw,
            flag_map=flag_map,
        )
        samhsa_id = fields.get("samhsa_id")
        key = _match_key(name, fields.get("city"), fields.get("state"))

        existing = None
        if samhsa_id and samhsa_id in existing_by_samhsa:
            existing = existing_by_samhsa[samhsa_id]
        elif key in existing_by_key:
            existing = existing_by_key[key]

        if existing:
            _apply_update(existing, fields, publish=publish)
            if samhsa_id:
                existing_by_samhsa[samhsa_id] = existing
            existing_by_key[key] = existing
            result.updated += 1
            continue

        pending = None
        if samhsa_id and samhsa_id in pending_by_samhsa:
            pending = pending_by_samhsa[samhsa_id]
        elif key in pending_by_key:
            pending = pending_by_key[key]
        if pending:
            pending.update({k: v for k, v in _new_mapping(fields, pending["slug"], publish=publish, now=now).items() if k != "slug"})
            result.updated += 1
            continue

        mapping = _new_mapping(fields, _next_slug(_slugify(name), taken_slugs), publish=publish, now=now)
        inserts.append(mapping)
        if samhsa_id:
            pending_by_samhsa[samhsa_id] = mapping
        pending_by_key[key] = mapping
        result.created += 1

        done = result.created + result.updated + result.skipped
        if done == 1 or done % 200 == 0 or done == planned:
            _emit(
                on_progress,
                phase="importing",
                message=f"Importing {done:,} of {planned:,} rows…",
                processed=done,
                total=planned,
                created=result.created,
                updated=result.updated,
                skipped=result.skipped,
                percent=_import_percent("importing", done, planned),
            )

    try:
        saved = 0
        insert_total = max(len(inserts), 1)
        _emit(
            on_progress,
            phase="saving",
            message="Saving listings…",
            processed=0,
            total=len(inserts),
            created=result.created,
            updated=result.updated,
            skipped=result.skipped,
            percent=85,
        )
        for batch in _chunked(inserts, INSERT_BATCH):
            db.bulk_insert_mappings(RehabCenter, batch)
            db.flush()
            saved += len(batch)
            _emit(
                on_progress,
                phase="saving",
                message=f"Saving listings… {saved:,} of {len(inserts):,}",
                processed=saved,
                total=len(inserts),
                created=result.created,
                updated=result.updated,
                skipped=result.skipped,
                percent=_import_percent("saving", saved, insert_total),
            )
        db.commit()
    except Exception as exc:  # noqa: BLE001
        db.rollback()
        result.errors = [f"Import failed while saving: {exc}"]
        result.created = 0
        result.updated = 0
        result.skipped = result.total_rows
        return result

    _emit(
        on_progress,
        phase="done",
        message="Import complete",
        processed=result.total_rows,
        total=result.total_rows,
        created=result.created,
        updated=result.updated,
        skipped=result.skipped,
        percent=100,
    )
    return result


def import_centers_csv(db: Session, content: str | bytes, *, publish: bool = True, filename: str = "import.csv") -> ImportResult:
    return import_centers_file(db, content, filename=filename, publish=publish)
